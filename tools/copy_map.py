import json,re,subprocess
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import os;B=os.path.dirname(os.path.abspath(__file__))+'/../'
D=json.loads(subprocess.check_output(['node','-e',"eval(require('fs').readFileSync('data.js','utf8').replace(/^const DATA/m,'global.DATA'));process.stdout.write(JSON.stringify(DATA))"],cwd=B))
cjk=re.compile(r'[\u4e00-\u9fff]')
rows=[]
def walk(v,p):
    if isinstance(v,list):
        if len(v)==2 and all(isinstance(x,str) for x in v) and cjk.search(v[0]) and not cjk.search(v[1]):
            rows.append((p,v[0],v[1]));return
        for i,x in enumerate(v): walk(x,f'{p}[{i}]')
    elif isinstance(v,dict):
        for k,x in v.items(): walk(x,f'{p}.{k}' if p else k)
    elif isinstance(v,str) and cjk.search(v): rows.append((p,v,''))
walk(D,'')
for _k,_n in D['nodes'].items():  # R4-4: English-only short headlines (small screens)
    for _h in 'UD':
        if _n[_h].get('h1s'): rows.append((f'nodes.{_k}.{_h}.h1s',_n[_h]['h1'][0],_n[_h]['h1s']))
# 2026-10-28i：index.html 內建的介面字（L?'英':'中'），路徑 index.html:<英文>；sheet_sync apply 只換中文那半
_IDX=open(B+'index.html',encoding='utf8').read()
for _en,_zh in dict.fromkeys(re.findall(r"L\?'([^']+)':'([^']*[\u4e00-\u9fff][^']*)'",_IDX)):
    rows.append((f'index.html:{_en}',_zh,_en))
legacy=[r'^nodes\.\w+\.(setup|chat|boss\.U)',r'^call\.',r'^callN\.',r'^self\.',r'^open2\.(reu|calling)',r'^open\[',r'^ui\.(reu|promise|intro|obj19|end|kids|baby)$',r'^nodes\.\w+\.src$',r'^v5\.blank$']
SEC={'wish':'2018 天燈願望','wish2':'新願望','wish26':'2026 願望','lantern':'結局天燈（玩家選老顏的燈）','v5':'介面','ui':'介面','open2':'開場','life':'老顏生活框','ring_off':'老顏生活框','verdict':'結局判詞','tag':'標籤','names':'人名','call':'電話（舊）','callN':'電話（舊）','self':'老顏（舊）','open':'開場（舊）','class_':'人名','babies':'其他','slug':'其他','sources':'結局出處頁（S7＋誰在倡議）','index.html':'介面（程式內建）'}
ORDER=['開場','2018 天燈願望','同學群組','逐年節點','選圖牆','群組後果（選上）','老顏生活框','新願望','2026 願望','結局天燈（玩家選老顏的燈）','結局判詞','結局私訊','結局出處頁（S7＋誰在倡議）','介面','介面（程式內建）','標籤','人名']
def classify(p):
    m=re.match(r'nodes\.y(\w+)\.(.*)',p)
    if m:
        y,rest=m.groups()
        f={'ask[0]':'主管交辦第1行','ask[1]':'主管交辦第2行（命題）','boss.U':'主管回饋（選上）','boss.D':'主管退件（選下，後果屏）'}
        if rest in f: return '逐年節點',y,f[rest],'主管'
        mm=re.match(r'([UD])\.(\w+)',rest)
        if mm:
            side='上半・文宣' if mm.group(1)=='U' else '下半・現場'
            ff={'who':'被拍的人','tag':'身分標籤','h1':'海報標題','h1s':'海報標題（英文小螢幕短版）','sub':'海報副標','aside':'現場照旁註'}.get(mm.group(2),mm.group(2))
            return '選圖牆',y,side+'・'+ff,''
        ms=re.match(r'setup\[(\d+)\]',rest)
        if ms: return '逐年節點',y,'舊前置對話',D['nodes']['y'+y]['setup'][int(ms.group(1))][0]
        return '逐年節點',y,rest,''
    m=re.match(r'group\.y(\w+)\[(\d+)\]',p)
    if m:
        g=D['group']['y'+m.group(1)][int(m.group(2))]
        return '同學群組',m.group(1),'第%d句'%(int(m.group(2))+1),g[0]
    m=re.match(r'gcons\.y(\w+)\.(\d)\[(\d+)\]',p)
    if m:
        g=D['gcons']['y'+m.group(1)][m.group(2)][int(m.group(3))]
        return '群組後果（選上）',m.group(1),'被犧牲者第%s次・第%d句%s'%(m.group(2),int(m.group(3))+1,('（'+g[1]+'）' if g[1] else '')),g[0]
    m=re.match(r'news\.y(\w+)\[(\d)\]',p)
    if m: return '群組後果（選上）',m.group(1),['新聞卡標題','新聞卡出處'][int(m.group(2))],''
    m=re.match(r'reunion\.([^|]+)\|([^.]+)\.(.*)',p)
    if m:
        f=m.group(3); f={'line':'私訊','off':'已讀不回／離線'}.get(f,f.replace('qa','追問'))
        kind={'gone':'退群版','stay':'受傷一次・還在群組版'}.get(D['reunion'][m.group(1)+'|'+m.group(2)].get('kind'),'')
        return '結局私訊',m.group(2).replace('y',''),kind+'・'+f,m.group(1)
    if p.startswith('index.html:'): return '介面（程式內建）','',p.split(':',1)[1],''
    k=re.split(r'[.\[]',p)[0]
    rest=p[len(k):].lstrip('.')
    spk=''
    m=re.match(r'^([^.\[|]+)',rest)
    if k in('wish','wish2','wish26','call','self','names','tag') and m: spk=m.group(1)
    if k=='callN': spk=''
    if k=='lantern':  # d29（2026-10-29）：老顏的燈限定選項 lantern.opt.<鍵>.(lamp|rep[.<回覆者>])，plamp＝八年後那盞
        ml=re.match(r'opt\.(\w+)\.?(lamp|rep)?\.?(.*)$',rest)
        if ml:
            g,f,who=ml.groups();lab={'s0':'①她自己・沒看過生活框','s1':'①她自己・生活框第1晚（冷飯）','s2':'①她自己・生活框第2晚（取卵再約）','s3':'①她自己・生活框第3晚（戒指沒人戴）','s4':'①她自己・生活框第4晚起（相框一個人）',
              'gone':'②朋友・已退群（{n}＝朋友名）','once':'②朋友・受傷一次','plamp':'③八年後','p0':'③八年後・沒人退群','p1':'③八年後・有人退群'}.get(g,g)
            if g=='plamp' or f=='lamp': return SEC[k],'',lab+'・燈上字＝按鈕字','老顏'
            return SEC[k],'',lab+'・回覆',(who or '佩珊')
        return SEC[k],'',rest,''
    return SEC.get(k,k),'',rest,spk
out=[]
for p,zh,en in rows:
    sec,y,field,spk=classify(p)
    dep=''
    m=re.match(r'group\.y(\w+)\[(\d+)\]',p)
    if m: dep='、'.join(D['group']['y'+m.group(1)][int(m.group(2))][2])
    m=re.match(r'gcons\.y(\w+)\.(\d)\[(\d+)\]\[2\]',p)
    if m: dep='、'.join(D['gcons']['y'+m.group(1)][m.group(2)][int(m.group(3))][3])
    if re.match(r'gcons\.[^\[]+\[\d+\]\[[013]',p): continue
    status='舊版未使用' if any(re.search(r,p) for r in legacy) else '使用中'
    if re.match(r'lantern\.(rules|names_en|prof_)|lantern\.opt\.\w+\.who',p): continue  # d29: who＝回覆者候選名單，不是文案
    if p.startswith('class_') or p.startswith('slug') or p.startswith('babies'): continue
    out.append([sec,y,field,spk,zh,en,dep,status,p])
yr=lambda y:(y or '0000')
out.sort(key=lambda r:(r[7]!='使用中',ORDER.index(r[0]) if r[0] in ORDER else 99,yr(r[1])))
wb=Workbook();ws=wb.active;ws.title='全部文案'
hdr=['編號','區塊','年份/節點','欄位','說話者','中文','English','依賴（此人退群就不出現）','狀態','資料路徑']
ws.append(hdr)
for i,r in enumerate(out,1): ws.append([i]+r)
for i,w in enumerate([6,16,10,22,8,46,46,14,10,30],1): ws.column_dimensions[get_column_letter(i)].width=w
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='C2185B')
for r in ws.iter_rows(min_row=2):
    for c in r: c.alignment=Alignment(wrap_text=True,vertical='top')
    if r[8].value!='使用中':
        for c in r: c.font=Font(color='999999')
ws.freeze_panes='F2';ws.auto_filter.ref=ws.dimensions
s2=wb.create_sheet('政策出處');s2.append(['節點','年份','出處'])
for k,v in D['nodes'].items(): s2.append([k[1:],v['year'],v['src']])
s2.column_dimensions['C'].width=90
# L：倡議者清單（每選項，含出處日期與連結）
import sys as _s;_s.path.insert(0,B+'tools');import sources_zh as _SZ
s4=wb.create_sheet('誰在倡議');s4.append(['節點','面','倡議者','摘要（中）','摘要（英）','出處 日期','連結'])
for a in _SZ.ADV: s4.append([a[1][1:],{'U':'上半・文宣','D':'下半・現場'}[a[2]],a[3],a[6],D['sources']['cards'][a[0]]['f'][0][1],a[4],a[5]])
for c,w in zip('ABCDEFG',[8,10,26,60,60,30,50]): s4.column_dimensions[c].width=w
s3=wb.create_sheet('說明')
for l in ['《我們與孩子的距離》Baby or Bust 全部文案與台詞對照（data.js，V=20261028i，2026-10-28 匯出）',
 '區塊依遊戲順序排列：開場 → 每年（同學群組 → 主管交辦 → 選圖牆 → 後果）→ 結局 → 介面。',
 '狀態「舊版未使用」＝資料檔還留著、遊戲不會顯示（灰字，排在最後）。',
 '依賴：列出的同學退出群組時，這句不出現；「!宜君」＝宜君退群後才出現。',
 '年份 2019a／2019b 等＝分支節點；結局私訊的年份欄＝該同學被犧牲的年份組合。',
 '要改文案：直接在表上改「中文」欄或加註解，告訴 Hark，由 Hark 回寫原始檔並重建。回寫只改中文，英文欄不重譯、不動（除非另外要求）；Hark 覆蓋這張表前會先讀回表上的中文修改並回寫，不會蓋掉。']: s3.append([l])
s3.column_dimensions['A'].width=110
wb.save(B+'qa/我們與孩子的距離_文案台詞對照.xlsx')
print(len(out),Counter(r[0] for r in out),Counter(r[7] for r in out))
